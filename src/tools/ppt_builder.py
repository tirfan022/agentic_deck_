from io import BytesIO
import json
import logging

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

logger = logging.getLogger(__name__)


def add_text(slide, text, x, y, w, h, size, bold=False, color=None):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.TOP

    paragraph = tf.paragraphs[0]
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color or RGBColor(35, 42, 55)
    return box


def normalize_bullets(raw_content):
    """Convert supported slide content formats into clean text bullets."""
    if isinstance(raw_content, str):
        bullets = raw_content.splitlines()
    elif isinstance(raw_content, list):
        bullets = raw_content
    elif isinstance(raw_content, dict):
        bullets = [
            value if isinstance(value, str) else str(value)
            for value in raw_content.values()
            if value is not None
        ]
    else:
        bullets = []

    clean_bullets = []
    for bullet in bullets:
        if isinstance(bullet, (dict, list)):
            bullet = json.dumps(bullet, ensure_ascii=False)
        text = str(bullet).strip()
        if text:
            clean_bullets.append(text.lstrip("-*•· ").strip())
    return clean_bullets


def build_presentation(slides_data, topic=""):
    """Build a text-only PowerPoint presentation."""
    if not isinstance(slides_data, list):
        raise ValueError("slides_data must be a list.")
    if len(slides_data) > 30:
        raise ValueError("A maximum of 30 slides is supported.")
    if not slides_data:
        raise ValueError("No slides were provided.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Cover slide
    cover = prs.slides.add_slide(blank)
    cover.background.fill.solid()
    cover.background.fill.fore_color.rgb = RGBColor(248, 250, 253)

    add_text(
        cover, "AGENTICDECK",
        0.8, 1.15, 10.8, 0.45,
        16, True, RGBColor(40, 85, 150),
    )
    add_text(
        cover, topic or "Your AI-generated presentation",
        0.8, 2.0, 11.5, 1.5,
        30, True, RGBColor(25, 35, 55),
    )
    add_text(
        cover, "Research-backed insights, structured for clarity.",
        0.85, 3.65, 10.8, 0.8,
        19, False, RGBColor(85, 95, 110),
    )

    for index, item in enumerate(slides_data, start=1):
        if not isinstance(item, dict):
            logger.warning("Skipping malformed slide %d.", index)
            continue

        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

        title = str(item.get("title") or f"Slide {index}").strip()
        add_text(
            slide, title[:180],
            0.65, 0.45, 12.0, 0.85,
            28, True, RGBColor(25, 43, 70),
        )

        rule = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.68), Inches(1.38),
            Inches(1.05), Inches(0.055),
        )
        rule.fill.solid()
        rule.fill.fore_color.rgb = RGBColor(45, 110, 190)
        rule.line.fill.background()

        bullets = normalize_bullets(item.get("content", []))
        body = slide.shapes.add_textbox(
            Inches(0.75), Inches(1.72),
            Inches(11.85), Inches(4.95),
        )
        tf = body.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = Inches(0.06)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.08)

        if not bullets:
            p = tf.paragraphs[0]
            p.text = "No slide content was provided."
            p.font.name = "Aptos"
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(150, 70, 70)
        else:
            for bullet_index, bullet in enumerate(bullets):
                p = tf.paragraphs[0] if bullet_index == 0 else tf.add_paragraph()
                p.text = bullet
                p.level = 0
                p.font.name = "Aptos"
                p.font.size = Pt(19)
                p.font.color.rgb = RGBColor(45, 52, 65)
                p.space_after = Pt(14)
                p.line_spacing = 1.1

        add_text(
            slide, f"AGENTICDECK · {index:02d}",
            0.72, 7.05, 5.0, 0.22,
            8, False, RGBColor(120, 128, 140),
        )

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output
